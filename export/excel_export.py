import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from models.schemas import NormalizationResult
import json

class ExcelExporter:
    def __init__(self, result: NormalizationResult):
        self.result = result
        
    def export(self, file_path: str) -> str:
        wb = openpyxl.Workbook()
        
        header_fill = PatternFill(start_color="1B4F72", end_color="1B4F72", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        violation_fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
        fixed_fill = PatternFill(start_color="CCFFCC", end_color="CCFFCC", fill_type="solid")
        pk_font = Font(color="008000", bold=True)
        fk_font = Font(color="FF8C00", bold=True)
        
        # Helper to style a sheet
        def style_sheet(ws):
            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = (max_length + 2)
                ws.column_dimensions[column].width = adjusted_width
        
        # 1. Original Data
        ws1 = wb.active
        ws1.title = "Original Data"
        if self.result.original_table:
            cols = [c.name for c in self.result.original_table.columns]
            ws1.append(cols)
            for cell in ws1[1]:
                cell.fill = header_fill
                cell.font = header_font
            
            for row_data in self.result.original_table.data:
                ws1.append([str(row_data.get(c, "")) for c in cols])
        style_sheet(ws1)
        
        # 2. 1NF Result, 3. 2NF Result, 4. 3NF Result
        for step_result in self.result.steps:
            nf = step_result.normal_form
            ws = wb.create_sheet(title=f"{nf} Result")
            current_row = 1
            
            for table in step_result.resulting_tables:
                ws.cell(row=current_row, column=1, value=f"Table: {table.name}")
                ws.cell(row=current_row, column=1).font = Font(bold=True)
                current_row += 1
                
                cols = [c.name for c in table.columns]
                for col_idx, col in enumerate(table.columns, 1):
                    cell = ws.cell(row=current_row, column=col_idx, value=col.name)
                    cell.fill = header_fill
                    if col.is_primary_key:
                        cell.font = pk_font
                    elif col.is_foreign_key:
                        cell.font = fk_font
                    else:
                        cell.font = header_font
                
                current_row += 1
                for row_data in table.data:
                    row_vals = [str(row_data.get(c, "")) for c in cols]
                    for col_idx, val in enumerate(row_vals, 1):
                        ws.cell(row=current_row, column=col_idx, value=val)
                    current_row += 1
                current_row += 2 # gap
            style_sheet(ws)
            
        # 5. SQL Queries
        ws_sql = wb.create_sheet(title="SQL Queries")
        ws_sql.append(["Step", "Query"])
        for cell in ws_sql[1]:
            cell.fill = header_fill
            cell.font = header_font
            
        for step in self.result.steps:
            for q in step.sql_queries:
                ws_sql.append([step.normal_form, q])
        style_sheet(ws_sql)
        
        # 6. Table Relationships
        ws_rel = wb.create_sheet(title="Table Relationships")
        ws_rel.append(["Table", "Column", "Is PK", "References (FK)"])
        for cell in ws_rel[1]:
            cell.fill = header_fill
            cell.font = header_font
            
        for table in self.result.final_tables:
            for col in table.columns:
                if col.is_primary_key or col.is_foreign_key:
                    ws_rel.append([table.name, col.name, "Yes" if col.is_primary_key else "No", col.references if col.is_foreign_key else ""])
        style_sheet(ws_rel)
        
        # 7. Summary
        ws_sum = wb.create_sheet(title="Summary")
        ws_sum.append(["Metric", "Value"])
        for cell in ws_sum[1]:
            cell.fill = header_fill
            cell.font = header_font
            
        total_violations = sum(len(s.violations_found) for s in self.result.steps)
        ws_sum.append(["Session ID", self.result.session_id])
        ws_sum.append(["Original Columns", len(self.result.original_table.columns)])
        ws_sum.append(["Final Tables", len(self.result.final_tables)])
        ws_sum.append(["Total Violations Found", total_violations])
        style_sheet(ws_sum)
        
        wb.save(file_path)
        return file_path
